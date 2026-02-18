#!/usr/bin/env python3
"""
Génération du Business Plan – Agriculture Hydroponique & Transformation
Kaffrine, Sénégal
Fichier de sortie : BUSINESS_PLAN_HYDROPONIE_KAFFRINE.docx
"""

import os
from docx import Document
from docx.shared import Pt, Cm, Inches, RGBColor, Emu
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.section import WD_ORIENT
from docx.oxml.ns import qn, nsdecls
from docx.oxml import parse_xml

OUTPUT = os.path.join(os.path.dirname(os.path.abspath(__file__)),
                      "BUSINESS_PLAN_HYDROPONIE_KAFFRINE.docx")

FONT_NAME = "Times New Roman"
RED = RGBColor(0xFF, 0x00, 0x00)
BLACK = RGBColor(0x00, 0x00, 0x00)
GREY_FILL = "D9D9D9"


# ─────────────────────────── helpers ───────────────────────────

def set_cell_shading(cell, color_hex):
    shading = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>')
    cell._tc.get_or_add_tcPr().append(shading)


def set_table_borders(table):
    tbl = table._tbl
    tblPr = tbl.tblPr if tbl.tblPr is not None else parse_xml(f'<w:tblPr {nsdecls("w")}/>')
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        '  <w:top w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        '  <w:left w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        '  <w:bottom w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        '  <w:right w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        '  <w:insideH w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        '  <w:insideV w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        '</w:tblBorders>'
    )
    tblPr.append(borders)


def run_text(paragraph, text, bold=False, size=Pt(12), color=BLACK, font=FONT_NAME):
    r = paragraph.add_run(text)
    r.bold = bold
    r.font.size = size
    r.font.color.rgb = color
    r.font.name = font
    rPr = r._element.get_or_add_rPr()
    rFonts = parse_xml(f'<w:rFonts {nsdecls("w")} w:ascii="{font}" w:hAnsi="{font}" w:eastAsia="{font}" w:cs="{font}"/>')
    rPr.insert(0, rFonts)
    return r


def placeholder(paragraph, text="XX", size=Pt(12)):
    return run_text(paragraph, text, bold=True, size=size, color=RED)


def add_paragraph(doc, text="", bold=False, size=Pt(12), alignment=WD_ALIGN_PARAGRAPH.JUSTIFY,
                  color=BLACK, space_after=Pt(6), space_before=Pt(0), first_line_indent=None):
    p = doc.add_paragraph()
    p.alignment = alignment
    pf = p.paragraph_format
    pf.space_after = space_after
    pf.space_before = space_before
    pf.line_spacing = Pt(18)
    if first_line_indent:
        pf.first_line_indent = first_line_indent
    if text:
        run_text(p, text, bold=bold, size=size, color=color)
    return p


def add_heading_custom(doc, text, level=1, size=Pt(14), alignment=WD_ALIGN_PARAGRAPH.LEFT):
    p = doc.add_paragraph()
    p.alignment = alignment
    pf = p.paragraph_format
    pf.space_before = Pt(12)
    pf.space_after = Pt(6)
    pf.line_spacing = Pt(18)
    pf.keep_with_next = True
    run_text(p, text, bold=True, size=size)
    return p


def add_subheading(doc, text, size=Pt(13)):
    return add_heading_custom(doc, text, level=2, size=size)


def add_sub2heading(doc, text, size=Pt(12)):
    return add_heading_custom(doc, text, level=3, size=size)


def add_bullet(doc, text, level=0, placeholder_parts=None):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    pf = p.paragraph_format
    pf.space_after = Pt(2)
    pf.line_spacing = Pt(18)
    pf.left_indent = Cm(1.27 * (level + 1))
    pf.first_line_indent = Cm(-0.63)
    run_text(p, "• ")
    if placeholder_parts:
        for part in placeholder_parts:
            if part[0] == "text":
                run_text(p, part[1])
            elif part[0] == "xx":
                placeholder(p, part[1] if len(part) > 1 else "XX")
    else:
        run_text(p, text)
    return p


def add_table(doc, headers, rows, col_widths=None, header_size=Pt(10), cell_size=Pt(10)):
    ncols = len(headers)
    table = doc.add_table(rows=1 + len(rows), cols=ncols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(table)

    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        set_cell_shading(cell, GREY_FILL)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if isinstance(h, tuple) and h[0] == "xx":
            placeholder(p, h[1] if len(h) > 1 else "XX", size=header_size)
        else:
            run_text(p, str(h), bold=True, size=header_size)

    for ri, row in enumerate(rows):
        for ci, val in enumerate(row):
            cell = table.rows[ri + 1].cells[ci]
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            if isinstance(val, tuple) and val[0] == "xx":
                placeholder(p, val[1] if len(val) > 1 else "XX", size=cell_size)
            else:
                run_text(p, str(val), size=cell_size)

    if col_widths:
        for ri, row in enumerate(table.rows):
            for ci, w in enumerate(col_widths):
                row.cells[ci].width = w

    return table


def add_page_break(doc):
    p = doc.add_paragraph()
    run = p.add_run()
    br = parse_xml(f'<w:br {nsdecls("w")} w:type="page"/>')
    run._element.append(br)
    return p


def fmt_money(val):
    if isinstance(val, (int, float)):
        s = f"{int(val):,}".replace(",", " ")
        return s
    return str(val)


# ─────────────────────── page numbering ────────────────────────

def add_page_numbers(doc):
    for section in doc.sections:
        footer = section.footer
        footer.is_linked_to_previous = False
        p = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        fldChar1 = parse_xml(f'<w:fldChar {nsdecls("w")} w:fldCharType="begin"/>')
        r1 = parse_xml(f'<w:r {nsdecls("w")}><w:rPr><w:rFonts w:ascii="{FONT_NAME}" w:hAnsi="{FONT_NAME}"/><w:sz w:val="20"/></w:rPr></w:r>')
        r1.append(fldChar1)
        instrText = parse_xml(f'<w:instrText {nsdecls("w")} xml:space="preserve"> PAGE </w:instrText>')
        r2 = parse_xml(f'<w:r {nsdecls("w")}><w:rPr><w:rFonts w:ascii="{FONT_NAME}" w:hAnsi="{FONT_NAME}"/><w:sz w:val="20"/></w:rPr></w:r>')
        r2.append(instrText)
        fldChar2 = parse_xml(f'<w:fldChar {nsdecls("w")} w:fldCharType="end"/>')
        r3 = parse_xml(f'<w:r {nsdecls("w")}><w:rPr><w:rFonts w:ascii="{FONT_NAME}" w:hAnsi="{FONT_NAME}"/><w:sz w:val="20"/></w:rPr></w:r>')
        r3.append(fldChar2)
        p._element.append(r1)
        p._element.append(r2)
        p._element.append(r3)


# ───────────────────── section builders ────────────────────────

def build_cover_page(doc):
    for _ in range(6):
        add_paragraph(doc, "", size=Pt(12))

    add_paragraph(doc, "BUSINESS PLAN", bold=True, size=Pt(18),
                  alignment=WD_ALIGN_PARAGRAPH.CENTER, space_after=Pt(24))

    add_paragraph(doc, "PROJET D'ENTREPRENEURIAT", bold=True, size=Pt(14),
                  alignment=WD_ALIGN_PARAGRAPH.CENTER, space_after=Pt(6))
    add_paragraph(doc, "AGRICULTURE HYDROPONIQUE & TRANSFORMATION", bold=True, size=Pt(14),
                  alignment=WD_ALIGN_PARAGRAPH.CENTER, space_after=Pt(36))

    p = add_paragraph(doc, "", alignment=WD_ALIGN_PARAGRAPH.CENTER, space_after=Pt(6))
    run_text(p, "Lieu : ", bold=True, size=Pt(12))
    run_text(p, "Kaffrine, Sénégal", size=Pt(12))

    p = add_paragraph(doc, "", alignment=WD_ALIGN_PARAGRAPH.CENTER, space_after=Pt(6))
    run_text(p, "Date : ", bold=True, size=Pt(12))
    run_text(p, "Février 2026", size=Pt(12))

    p = add_paragraph(doc, "", alignment=WD_ALIGN_PARAGRAPH.CENTER, space_after=Pt(6))
    run_text(p, "Promoteur(s) : ", bold=True, size=Pt(12))
    placeholder(p)

    add_page_break(doc)


def build_toc(doc):
    add_heading_custom(doc, "TABLE DES MATIÈRES", size=Pt(14),
                       alignment=WD_ALIGN_PARAGRAPH.CENTER)
    add_paragraph(doc, "")

    toc_items = [
        "INTRODUCTION",
        "",
        "1ère PARTIE : L'ENTREPRISE ET LES PROMOTEURS",
        "    I. Présentation générale du projet",
        "    II. Objectifs de croissance",
        "    III. Vision et Mission",
        "    IV. Présentation des promoteurs",
        "    V. Cadre juridique et fiscal",
        "",
        "2ème PARTIE : L'ÉTUDE DU MARCHÉ",
        "    I. Analyse du marché",
        "        1- Analyse de la demande",
        "        2- Analyse de l'offre",
        "    II. Analyse de l'environnement",
        "    III. Business Model Canvas",
        "    IV. Prévisions des ventes",
        "",
        "3ème PARTIE : ÉTUDE TECHNIQUE ET OPÉRATIONNELLE",
        "    I. Cadre technique",
        "    II. Cadre organisationnel",
        "",
        "4ème PARTIE : ÉTUDE FINANCIÈRE",
        "    I. Investissements et coût total",
        "    II. Plan de financement initial",
        "    III. Remboursement de l'emprunt",
        "    IV. Amortissement des immobilisations",
        "    V. Compte d'exploitation prévisionnel",
        "    VI. Valeur Actuelle Nette (VAN)",
        "    VII. Seuil de rentabilité",
        "    VIII. Plan de financement sur 5 ans",
        "",
        "5ème PARTIE : IMPACTS SOCIO-ÉCONOMIQUES ET GESTION DES RISQUES",
        "    I. Impact économique",
        "    II. Impact social",
        "    III. Impact environnemental",
        "    IV. Risques et mesures d'atténuation",
        "",
        "CONCLUSION",
        "",
        "ANNEXES",
        "    Annexe 1 : Questionnaire d'enquête",
        "    Annexe 2 : Tableau des produits et prix",
    ]

    for item in toc_items:
        if item == "":
            add_paragraph(doc, "", space_after=Pt(2))
        else:
            indent = 0
            clean = item
            while clean.startswith("    "):
                indent += 1
                clean = clean[4:]
            p = add_paragraph(doc, "", space_after=Pt(2))
            p.paragraph_format.left_indent = Cm(1.0 * indent)
            bld = indent == 0 and clean not in ["INTRODUCTION", "CONCLUSION"]
            if clean.startswith("Annexe"):
                bld = False
            run_text(p, clean, bold=bld, size=Pt(12))

    add_page_break(doc)


def build_introduction(doc):
    add_heading_custom(doc, "INTRODUCTION", size=Pt(14),
                       alignment=WD_ALIGN_PARAGRAPH.CENTER)
    add_paragraph(doc, "")
    text = (
        "Dans un contexte marqué par une forte demande en produits alimentaires sains et locaux, "
        "le présent projet vise la création d'une entreprise agroalimentaire moderne basée sur la "
        "culture hors sol (hydroponie), combinée à une unité artisanale de transformation. "
        "L'entreprise aura pour activité principale la production de fruits et légumes à haute valeur "
        "ajoutée, ainsi que leur transformation en produits dérivés (jus, compotes, confitures, sauces "
        "et tartinades). Les produits seront commercialisés localement sous une marque propre."
    )
    add_paragraph(doc, text, first_line_indent=Cm(1.27))

    text2 = (
        "Ce projet ambitionne de contribuer à la sécurité alimentaire, à la création d'emplois locaux "
        "et à la valorisation des produits agricoles de la région de Kaffrine. Le coût total du projet "
        "est estimé à 13 000 000 FCFA, financé à hauteur de 30 % par apport personnel et 70 % par "
        "emprunt bancaire."
    )
    add_paragraph(doc, text2, first_line_indent=Cm(1.27))

    add_page_break(doc)


# ──────────────── 1ère PARTIE ────────────────

def build_partie1(doc):
    add_heading_custom(doc, "1ère PARTIE : L'ENTREPRISE ET LES PROMOTEURS",
                       size=Pt(14), alignment=WD_ALIGN_PARAGRAPH.CENTER)
    add_paragraph(doc, "")

    # I. Présentation générale
    add_subheading(doc, "I. Présentation générale du projet")

    add_sub2heading(doc, "Historique")
    add_paragraph(doc, (
        "Ce projet est né du constat de la difficulté d'accès à des produits frais et transformés "
        "de qualité dans la région de Kaffrine, ainsi que du manque d'unités agricoles modernes. "
        "L'idée est venue de combiner l'agriculture innovante (hydroponie) avec la transformation "
        "locale afin d'augmenter la valeur des produits."
    ), first_line_indent=Cm(1.27))

    add_sub2heading(doc, "Opportunités")
    add_bullet(doc, "Demande croissante en produits naturels et sains")
    add_bullet(doc, "Faible concurrence structurée dans la région")
    add_bullet(doc, "Marchés locaux disponibles (marchés hebdomadaires, restaurants, écoles)")
    add_bullet(doc, "Politique nationale favorable (PSE, PRACAS)")

    add_sub2heading(doc, "Description du projet")
    add_bullet(doc, "Produire des légumes (tomate, laitue, poivron, concombre, aubergine) par culture hydroponique sous serre")
    add_bullet(doc, "Produire ou acheter localement certains fruits (fraise, melon, pastèque, papaye, mangue)")
    add_bullet(doc, "Transformer une partie de la production en jus, compotes, confitures, sauces et tartinades")
    add_bullet(doc, "Commercialiser les produits sous une marque locale")

    add_sub2heading(doc, "Mode d'utilisation")
    add_paragraph(doc, (
        "Produits frais pour consommation directe. Produits transformés conditionnés en "
        "bouteilles et pots."
    ), first_line_indent=Cm(1.27))

    add_sub2heading(doc, "Secteur d'activité")
    add_paragraph(doc, (
        "Agroalimentaire (agriculture moderne – hydroponie + transformation artisanale + "
        "commerce alimentaire)."
    ), first_line_indent=Cm(1.27))

    add_sub2heading(doc, "Points forts")
    add_bullet(doc, "Utilisation rationnelle de l'eau (goutte-à-goutte, 90 % d'économie)")
    add_bullet(doc, "Production sous serre (qualité, protection climatique)")
    add_bullet(doc, "Double activité : production + transformation")
    add_bullet(doc, "Marché local favorable")
    add_bullet(doc, "Production toute l'année")

    add_sub2heading(doc, "Points faibles")
    add_bullet(doc, "Investissement initial élevé (13 000 000 FCFA)")
    add_bullet(doc, "Besoin de formation technique")
    add_bullet(doc, "Dépendance à l'électricité et à l'eau")
    add_bullet(doc, "Technologie peu connue localement")

    add_sub2heading(doc, "Contraintes")
    add_bullet(doc, "Techniques : maîtrise de l'hydroponie, maintenance des équipements")
    add_bullet(doc, "Financières : coût des équipements, accès au crédit")
    add_bullet(doc, "Marché : pouvoir d'achat local limité")
    add_bullet(doc, "Réglementaires : autorisation sanitaire, RCCM, fiscalité")

    # II. Objectifs de croissance
    add_subheading(doc, "II. Objectifs de croissance")

    add_sub2heading(doc, "Moyen terme (1-3 ans)")
    add_bullet(doc, "Atteindre l'équilibre financier en moins de 24 mois")
    add_bullet(doc, "Transformer 1 tonne de produits par mois")
    add_bullet(doc, "Créer 10 à 15 emplois directs")
    add_bullet(doc, "Atteindre un chiffre d'affaires de 37 500 000 FCFA en année 3")
    add_bullet(doc, "Assurer une distribution régionale")

    add_sub2heading(doc, "Long terme (3-5 ans et au-delà)")
    add_bullet(doc, "Développer l'aquaponie (combinaison élevage poisson + hydroponie)")
    add_bullet(doc, "Construire de nouvelles serres")
    add_bullet(doc, "Lancer de nouveaux produits transformés")
    add_bullet(doc, "Accéder au marché national")
    add_bullet(doc, "Passer de SARL à SA")
    add_bullet(doc, "Atteindre un chiffre d'affaires de 56 000 000 FCFA en année 5")

    add_sub2heading(doc, "Tableau de progression")
    add_table(doc,
              ["Indicateur", "Année 1", "Année 2", "Année 3", "Année 4", "Année 5"],
              [
                  ["CA (FCFA)", "24 000 000", "30 000 000", "37 500 000", "46 000 000", "56 000 000"],
                  ["Emplois", "7", "9", "12", "14", "17"],
                  ["Forme juridique", "SARL", "SARL", "SARL", "SARL", "SA (objectif)"],
                  ["Zone distribution", "Kaffrine", "Kaffrine + environs", "Région", "Régional élargi", "National"],
              ])

    # III. Vision et Mission
    add_subheading(doc, "III. Vision et Mission")

    add_sub2heading(doc, "Vision")
    add_paragraph(doc, (
        "Devenir une référence régionale dans la production et la transformation de produits "
        "agricoles locaux, sains et durables."
    ), first_line_indent=Cm(1.27))

    add_sub2heading(doc, "Missions")
    add_bullet(doc, "Produire des aliments de qualité, accessibles et nutritifs")
    add_bullet(doc, "Valoriser les produits locaux par la transformation artisanale")
    add_bullet(doc, "Créer de l'emploi durable pour les jeunes et les femmes")
    add_bullet(doc, "Contribuer au développement économique local et à la sécurité alimentaire")

    # IV. Présentation des promoteurs
    add_subheading(doc, "IV. Présentation des promoteurs")

    add_sub2heading(doc, "Tableau des promoteurs")
    add_table(doc,
              ["Prénoms", "Nom", "Situation matrimoniale", "Diplômes / Formation", "Expériences"],
              [
                  [("xx", "XX"), ("xx", "XX"), ("xx", "XX"), ("xx", "XX"), ("xx", "XX")],
                  [("xx", "XX"), ("xx", "XX"), ("xx", "XX"), ("xx", "XX"), ("xx", "XX")],
              ])

    add_paragraph(doc, "")
    p = add_paragraph(doc, "")
    run_text(p, "Personnalité : ", bold=True)
    placeholder(p)

    p = add_paragraph(doc, "")
    run_text(p, "Potentiel et facteurs déterminants : ", bold=True)
    placeholder(p)

    add_sub2heading(doc, "Motivations")
    add_bullet(doc, "Autonomie financière")
    add_bullet(doc, "Création d'emplois pour les jeunes et les femmes")
    add_bullet(doc, "Développement économique local")
    add_bullet(doc, "Passion pour l'agriculture innovante")

    add_sub2heading(doc, "Actions correctives")
    add_paragraph(doc, (
        "Formations en hydroponie et transformation agroalimentaire auprès de l'ISRA, "
        "de l'ANCAR et des centres de formation professionnelle."
    ), first_line_indent=Cm(1.27))

    # V. Cadre juridique et fiscal
    add_subheading(doc, "V. Cadre juridique et fiscal")

    add_sub2heading(doc, "1- Cadre juridique")
    add_bullet(doc, "Forme juridique : SARL (Société à Responsabilité Limitée)")
    p = add_paragraph(doc, "")
    p.paragraph_format.left_indent = Cm(1.27)
    run_text(p, "• Dénomination sociale : ")
    placeholder(p)

    add_bullet(doc, "Objet social : Production, transformation et commercialisation de produits agroalimentaires")
    add_bullet(doc, "Siège social : Kaffrine, Sénégal")
    add_bullet(doc, "Durée : 99 ans")
    add_bullet(doc, "Capital social : 3 900 000 FCFA")
    add_bullet(doc, "Gouvernance : Gérant(s) + Assemblée Générale des associés")
    add_bullet(doc, "Frais de notaire estimés : ~ 300 000 FCFA")

    add_sub2heading(doc, "2- Cadre fiscal")
    add_paragraph(doc, (
        "En tant que SARL, l'entreprise est soumise aux obligations fiscales suivantes :"
    ))
    add_table(doc,
              ["Impôt / Taxe", "Base / Taux"],
              [
                  ["Impôt sur les Sociétés (IS)", "30 % du résultat fiscal"],
                  ["Patente", "Variable selon le CA"],
                  ["TVA", "18 %"],
                  ["CFCE (Contribution Forfaitaire à la Charge de l'Employeur)", "3 % de la masse salariale"],
                  ["IMF (Impôt Minimum Forfaitaire)", "0,5 % du CA (minimum 500 000 FCFA)"],
                  ["Droit d'enregistrement du bail", "Selon le montant du bail"],
              ])

    add_paragraph(doc, "")
    add_sub2heading(doc, "Justification du choix de la SARL")
    add_bullet(doc, "Responsabilité limitée aux apports des associés")
    add_bullet(doc, "Accès facilité aux financements bancaires")
    add_bullet(doc, "Gouvernance adaptée aux PME")
    add_bullet(doc, "Possibilité de transition vers SA en phase de croissance")

    add_page_break(doc)


# ──────────────── 2ème PARTIE ────────────────

def build_partie2(doc):
    add_heading_custom(doc, "2ème PARTIE : L'ÉTUDE DU MARCHÉ",
                       size=Pt(14), alignment=WD_ALIGN_PARAGRAPH.CENTER)
    add_paragraph(doc, "")

    # I. Analyse du marché
    add_subheading(doc, "I. Analyse du marché")

    # 1- Demande
    add_sub2heading(doc, "1- Analyse de la demande")

    p = add_paragraph(doc, "")
    run_text(p, "Cible : ", bold=True)
    run_text(p, "Ménages, restaurants, hôtels, cantines scolaires, boutiques alimentaires.")

    p = add_paragraph(doc, "")
    run_text(p, "Zone ciblée : ", bold=True)
    run_text(p, "Kaffrine et communes environnantes (~600 000 habitants pour la région).")

    p = add_paragraph(doc, "")
    run_text(p, "Typologie de la demande : ", bold=True)
    run_text(p, "Demande primaire (produits frais) + demande complémentaire (produits transformés).")

    p = add_paragraph(doc, "")
    run_text(p, "Taille du marché : ", bold=True)
    run_text(p, "Croissance de 5 à 8 % par an dans le secteur agroalimentaire au Sénégal. "
             "La région de Kaffrine est sous-desservie en produits frais et transformés de qualité.")

    add_sub2heading(doc, "a) Comportement du consommateur")
    add_paragraph(doc, (
        "Attentes : produits frais, de qualité, à prix abordable, sans produits chimiques, "
        "disponibles toute l'année."
    ), first_line_indent=Cm(1.27))

    p = add_paragraph(doc, "")
    run_text(p, "Fréquence d'achat : ", bold=True)
    placeholder(p)

    p = add_paragraph(doc, "")
    run_text(p, "Volume moyen : ", bold=True)
    placeholder(p)

    p = add_paragraph(doc, "")
    run_text(p, "Budget moyen : ", bold=True)
    placeholder(p)

    add_sub2heading(doc, "Enquête terrain")
    p = add_paragraph(doc, "")
    run_text(p, "Méthode : ", bold=True)
    run_text(p, "Échantillonnage aléatoire stratifié.")
    p = add_paragraph(doc, "")
    run_text(p, "Taille de l'échantillon : ", bold=True)
    placeholder(p)
    run_text(p, " personnes (minimum 100 recommandé).")

    add_paragraph(doc, (
        "Le questionnaire d'enquête est présenté en Annexe 1."
    ), first_line_indent=Cm(1.27))

    add_sub2heading(doc, "b) Analyse des données de l'enquête")
    p = add_paragraph(doc, "")
    placeholder(p, "XX — Résultats à compléter après enquête terrain.")
    add_paragraph(doc, "")

    p = add_paragraph(doc, "")
    run_text(p, "Diagramme 1 — Fréquence d'achat de légumes frais : ", bold=True)
    placeholder(p, "XX (insérer diagramme)")
    add_paragraph(doc, "")
    p = add_paragraph(doc, "")
    run_text(p, "Diagramme 2 — Budget hebdomadaire en légumes : ", bold=True)
    placeholder(p, "XX (insérer diagramme)")
    add_paragraph(doc, "")
    p = add_paragraph(doc, "")
    run_text(p, "Diagramme 3 — Intérêt pour les produits sous serre : ", bold=True)
    placeholder(p, "XX (insérer diagramme)")
    add_paragraph(doc, "")
    p = add_paragraph(doc, "")
    run_text(p, "Diagramme 4 — Prix acceptable pour un jus naturel local 33cl : ", bold=True)
    placeholder(p, "XX (insérer diagramme)")

    # 2- Offre
    add_sub2heading(doc, "2- Analyse de l'offre")

    p = add_paragraph(doc, "")
    run_text(p, "Concurrents directs : ", bold=True)
    placeholder(p)

    p = add_paragraph(doc, "")
    run_text(p, "Concurrents indirects : ", bold=True)
    placeholder(p)

    add_paragraph(doc, "")
    add_sub2heading(doc, "Tableau comparatif des 4P")
    add_table(doc,
              ["", "Place", "Produits", "Prix", "Promotion"],
              [
                  ["Concurrents directs", ("xx", "XX"), ("xx", "XX"), ("xx", "XX"), ("xx", "XX")],
                  ["Concurrents indirects", ("xx", "XX"), ("xx", "XX"), ("xx", "XX"), ("xx", "XX")],
                  ["Notre positionnement", ("xx", "XX"), ("xx", "XX"), ("xx", "XX"), ("xx", "XX")],
              ])

    # II. Analyse de l'environnement
    add_paragraph(doc, "")
    add_subheading(doc, "II. Analyse de l'environnement")

    add_sub2heading(doc, "Analyse PESTEL")
    add_table(doc,
              ["Facteur", "Éléments clés"],
              [
                  ["Politique", "Stabilité politique ; PSE (Plan Sénégal Émergent) ; PRACAS ; PNASAR ; DER/FJ ; FONGIP"],
                  ["Économique", "PIB +6-9 % ; inflation <3 % ; taux directeur BCEAO 3,25 % ; PME = 51 % crédits UEMOA ; pouvoir d'achat limité en zone rurale"],
                  ["Social", "Population jeune (60 % < 25 ans) ; urbanisation croissante ; conscience nutritionnelle ; chômage des jeunes"],
                  ["Technologique", "Technologies agricoles innovantes ; énergie solaire ; digitalisation (mobile money, réseaux sociaux)"],
                  ["Environnemental", "Changement climatique ; stress hydrique ; hydroponie = -90 % eau ; gestion des déchets organiques"],
                  ["Légal", "OHADA ; CGI Sénégal ; normes sanitaires ; autorisation de mise sur le marché ; convention collective"],
              ])

    add_paragraph(doc, "")
    add_sub2heading(doc, "Analyse SWOT")

    swot_table = doc.add_table(rows=3, cols=2)
    swot_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(swot_table)

    h1 = swot_table.rows[0].cells[0]
    set_cell_shading(h1, GREY_FILL)
    p = h1.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_text(p, "FORCES", bold=True, size=Pt(10))

    h2 = swot_table.rows[0].cells[1]
    set_cell_shading(h2, GREY_FILL)
    p = h2.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_text(p, "FAIBLESSES", bold=True, size=Pt(10))

    c = swot_table.rows[1].cells[0]
    p = c.paragraphs[0]
    run_text(p, "• Innovation hydroponie\n• Double activité (production + transformation)\n"
             "• Économie d'eau (90 %)\n• Production annuelle continue\n• Marque locale", size=Pt(10))

    c = swot_table.rows[1].cells[1]
    p = c.paragraphs[0]
    run_text(p, "• Investissement élevé (13 M FCFA)\n• Dépendance à l'électricité\n"
             "• Formation technique nécessaire\n• Faible notoriété initiale", size=Pt(10))

    h3 = swot_table.rows[2].cells[0]
    set_cell_shading(h3, GREY_FILL)
    p = h3.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_text(p, "OPPORTUNITÉS", bold=True, size=Pt(10))

    h4 = swot_table.rows[2].cells[1]
    set_cell_shading(h4, GREY_FILL)
    p = h4.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_text(p, "MENACES", bold=True, size=Pt(10))

    swot_table.add_row()
    c = swot_table.rows[3].cells[0]
    p = c.paragraphs[0]
    run_text(p, "• Demande croissante en produits sains\n• Faible concurrence structurée\n"
             "• Soutien étatique (DER, FONGIP)\n• Expansion régionale et nationale possible", size=Pt(10))

    c = swot_table.rows[3].cells[1]
    p = c.paragraphs[0]
    run_text(p, "• Pouvoir d'achat limité\n• Coupures d'électricité fréquentes\n"
             "• Concurrence des produits importés\n• Chaleur extrême", size=Pt(10))

    # III. Business Model Canvas
    add_paragraph(doc, "")
    add_subheading(doc, "III. Business Model Canvas")

    bmc_headers = ["Bloc", "Contenu"]
    bmc_rows = [
        ["Partenaires clés", "ISRA, ANCAR, fournisseurs intrants, banques/IMF, collectivités locales, distributeurs"],
        ["Activités clés", "Production hydroponique sous serre ; transformation artisanale ; commercialisation ; marketing"],
        ["Ressources clés", "Serres, système hydroponique, unité de transformation, personnel qualifié, fonds de roulement"],
        ["Proposition de valeur", "Produits frais et transformés locaux, sains, sans pesticides, disponibles toute l'année, à prix accessibles"],
        ["Relations clients", "Vente directe, fidélisation, service après-vente, réseaux sociaux, dégustations"],
        ["Canaux de distribution", "Point de vente propre, marchés hebdomadaires, livraison restaurants/hôtels, boutiques partenaires"],
        ["Segments clients", "Ménages, restaurants, hôtels, cantines scolaires, boutiques alimentaires, grossistes"],
        ["Structure des coûts", "Investissement initial 13 000 000 FCFA ; charges variables (intrants, emballages) ; charges fixes (salaires, loyer, énergie)"],
        ["Sources de revenus", "Vente de légumes frais ; vente de produits transformés (jus, confitures, sauces, compotes, tartinades)"],
    ]
    add_table(doc, bmc_headers, bmc_rows)

    # IV. Prévisions des ventes
    add_paragraph(doc, "")
    add_subheading(doc, "IV. Prévisions des ventes")
    add_table(doc,
              ["", "Année 1", "Année 2", "Année 3", "Année 4", "Année 5"],
              [
                  ["CA (FCFA)", "24 000 000", "30 000 000", "37 500 000", "46 000 000", "56 000 000"],
                  ["Croissance", "—", "+25 %", "+25 %", "+22,7 %", "+21,7 %"],
              ])

    add_page_break(doc)


# ──────────────── 3ème PARTIE ────────────────

def build_partie3(doc):
    add_heading_custom(doc, "3ème PARTIE : ÉTUDE TECHNIQUE ET OPÉRATIONNELLE",
                       size=Pt(14), alignment=WD_ALIGN_PARAGRAPH.CENTER)
    add_paragraph(doc, "")

    add_subheading(doc, "I. Cadre technique")

    add_sub2heading(doc, "Processus de production hydroponique")
    steps_hydro = [
        "Préparation des semis",
        "Transplantation dans le système hydroponique",
        "Gestion de la solution nutritive",
        "Croissance sous serre contrôlée",
        "Récolte",
        "Tri et conditionnement",
        "Stockage en chambre froide",
        "Commercialisation",
    ]
    flow = "  →  ".join([f"({i+1}) {s}" for i, s in enumerate(steps_hydro)])
    add_paragraph(doc, flow, size=Pt(11))

    add_sub2heading(doc, "Processus de transformation")
    steps_transfo = [
        "Réception et lavage des matières premières",
        "Épluchage et découpe",
        "Cuisson / Mixage / Pressage",
        "Ajout d'ingrédients (sucre, épices, conservateurs naturels)",
        "Pasteurisation",
        "Conditionnement (bouteilles, pots)",
        "Étiquetage",
        "Stockage",
        "Distribution",
    ]
    flow2 = "  →  ".join([f"({i+1}) {s}" for i, s in enumerate(steps_transfo)])
    add_paragraph(doc, flow2, size=Pt(11))

    add_sub2heading(doc, "Commentaire technique")
    add_paragraph(doc, (
        "Le système hydroponique retenu est de type NFT (Nutrient Film Technique) ou goutte-à-goutte "
        "sur substrat (fibre de coco ou perlite). La solution nutritive est recyclée en circuit fermé, "
        "ce qui permet une économie d'eau de l'ordre de 90 % par rapport à l'agriculture traditionnelle."
    ), first_line_indent=Cm(1.27))
    add_paragraph(doc, (
        "L'unité de transformation artisanale applique un protocole HACCP simplifié pour garantir "
        "la qualité sanitaire des produits. La pasteurisation assure une conservation de 3 à 6 mois "
        "pour les produits transformés."
    ), first_line_indent=Cm(1.27))

    # II. Cadre organisationnel
    add_subheading(doc, "II. Cadre organisationnel")

    add_sub2heading(doc, "Organigramme hiérarchique")
    add_paragraph(doc, "")
    p = add_paragraph(doc, "", alignment=WD_ALIGN_PARAGRAPH.CENTER)
    run_text(p, "Directeur Général / Gérant", bold=True, size=Pt(11))
    p = add_paragraph(doc, "", alignment=WD_ALIGN_PARAGRAPH.CENTER)
    run_text(p, "│", size=Pt(11))
    p = add_paragraph(doc, "", alignment=WD_ALIGN_PARAGRAPH.CENTER)
    run_text(p, "┌──────────────────────┼──────────────────────┐", size=Pt(11))
    p = add_paragraph(doc, "", alignment=WD_ALIGN_PARAGRAPH.CENTER)
    run_text(p, "│                                      │                                      │", size=Pt(10))
    p = add_paragraph(doc, "", alignment=WD_ALIGN_PARAGRAPH.CENTER)
    run_text(p, "Département                    Département                    Département", bold=True, size=Pt(10))
    p = add_paragraph(doc, "", alignment=WD_ALIGN_PARAGRAPH.CENTER)
    run_text(p, "Admin & Finance           Production & Technique       Commercial & Marketing", size=Pt(10))
    p = add_paragraph(doc, "", alignment=WD_ALIGN_PARAGRAPH.CENTER)
    run_text(p, "│                                      │                                      │", size=Pt(10))
    p = add_paragraph(doc, "", alignment=WD_ALIGN_PARAGRAPH.CENTER)
    run_text(p, "Agent admin.                 Ouvriers agricoles            Agent commercial", size=Pt(10))
    p = add_paragraph(doc, "", alignment=WD_ALIGN_PARAGRAPH.CENTER)
    run_text(p, "                                  Ouvriers transformation", size=Pt(10))

    add_paragraph(doc, "")
    add_sub2heading(doc, "Tableau des effectifs")
    add_table(doc,
              ["Poste", "Nombre", "Salaire mensuel (FCFA)", "Salaire annuel (FCFA)"],
              [
                  ["Gérant / DG", ("xx", "XX"), ("xx", "XX"), ("xx", "XX")],
                  ["Technicien hydroponique", ("xx", "XX"), ("xx", "XX"), ("xx", "XX")],
                  ["Agent administratif", ("xx", "XX"), ("xx", "XX"), ("xx", "XX")],
                  ["Agent commercial", ("xx", "XX"), ("xx", "XX"), ("xx", "XX")],
                  ["Ouvriers agricoles", ("xx", "XX"), ("xx", "XX"), ("xx", "XX")],
                  ["Ouvriers transformation", ("xx", "XX"), ("xx", "XX"), ("xx", "XX")],
                  ["Total", ("xx", "XX"), "—", ("xx", "XX")],
              ])

    add_page_break(doc)


# ──────────────── 4ème PARTIE ────────────────

def build_partie4(doc):
    add_heading_custom(doc, "4ème PARTIE : ÉTUDE FINANCIÈRE",
                       size=Pt(14), alignment=WD_ALIGN_PARAGRAPH.CENTER)
    add_paragraph(doc, "")

    add_paragraph(doc, (
        "Note : Les chiffres présentés ci-dessous sont des estimations prévisionnelles. "
        "Les postes marqués XX sont à compléter par le(s) promoteur(s) sur la base de devis réels."
    ), bold=True, size=Pt(11), first_line_indent=Cm(1.27))

    # I. Investissements
    add_subheading(doc, "I. Investissements et coût total du projet")

    add_table(doc,
              ["Désignations", "Montants (FCFA)"],
              [
                  ["Frais de constitution (notaire, RCCM, NINEA)", "300 000"],
                  ["Terrain + installations", ("xx", "XX")],
                  ["Bâtiments / Serres", ("xx", "XX")],
                  ["Machines (hydroponie, pompes, réservoirs)", ("xx", "XX")],
                  ["Dépôt et cautionnement (SENELEC, SDE, Téléphone)", ("xx", "XX")],
                  ["Matériels informatiques", ("xx", "XX")],
                  ["Matériels de télécommunication", ("xx", "XX")],
                  ["Matériels d'exploitation (transformation)", ("xx", "XX")],
                  ["Mobilier de bureau", ("xx", "XX")],
                  ["Matériel de transport", ("xx", "XX")],
                  ["Autres matériels", ("xx", "XX")],
                  ["Sous-total immobilisations", ("xx", "XX")],
                  ["Imprévus (15 %)", ("xx", "XX")],
                  ["Total investissements (I)", "10 650 000"],
                  ["BFR (Besoin en Fonds de Roulement — 3 mois)", "2 350 000"],
                  ["Coût total du projet = I + BFR", "13 000 000"],
              ])

    # II. Plan de financement initial
    add_paragraph(doc, "")
    add_subheading(doc, "II. Plan de financement initial")

    add_table(doc,
              ["Emplois", "Montants (FCFA)", "Ressources", "Montants (FCFA)"],
              [
                  ["Total investissements", "10 650 000", "Apport personnel (30 %)", "3 900 000"],
                  ["BFR (3 mois)", "2 350 000", "Emprunt bancaire (70 %)", "9 100 000"],
                  ["Coût total", "13 000 000", "Total ressources", "13 000 000"],
              ])

    # III. Remboursement emprunt
    add_paragraph(doc, "")
    add_subheading(doc, "III. Tableau de remboursement de l'emprunt")

    add_paragraph(doc, (
        "Montant emprunté : 9 100 000 FCFA — Taux d'intérêt : 10 % par an — Durée : 5 ans — "
        "Amortissement constant."
    ), bold=True, size=Pt(11))

    add_table(doc,
              ["Année", "Capital restant dû (début)", "Échéance principale (EP)", "Intérêts payés (IP)", "Annuité totale (EP + IP)"],
              [
                  ["1", "9 100 000", "1 820 000", "910 000", "2 730 000"],
                  ["2", "7 280 000", "1 820 000", "728 000", "2 548 000"],
                  ["3", "5 460 000", "1 820 000", "546 000", "2 366 000"],
                  ["4", "3 640 000", "1 820 000", "364 000", "2 184 000"],
                  ["5", "1 820 000", "1 820 000", "182 000", "2 002 000"],
                  ["Total", "—", "9 100 000", "2 730 000", "11 830 000"],
              ])

    # IV. Amortissement
    add_paragraph(doc, "")
    add_subheading(doc, "IV. Amortissement des immobilisations")

    add_paragraph(doc, (
        "Méthode retenue : amortissement linéaire. "
        "Formule : Dotation annuelle = Montant / Durée d'utilisation prévisionnelle (DUP)."
    ), size=Pt(11), first_line_indent=Cm(1.27))

    add_table(doc,
              ["Immobilisation", "Montant (FCFA)", "DUP (ans)", "Dotation annuelle (FCFA)"],
              [
                  ["Frais de constitution", "300 000", "5", "60 000"],
                  ["Terrain", "Non amortissable", "—", "—"],
                  ["Bâtiments / Serres", ("xx", "XX"), "5", ("xx", "XX")],
                  ["Machines (hydroponie, pompes, réservoirs)", ("xx", "XX"), "4", ("xx", "XX")],
                  ["Matériels informatiques", ("xx", "XX"), "3", ("xx", "XX")],
                  ["Matériels de télécommunication", ("xx", "XX"), "5", ("xx", "XX")],
                  ["Matériels d'exploitation (transformation)", ("xx", "XX"), "4", ("xx", "XX")],
                  ["Mobilier de bureau", ("xx", "XX"), "5", ("xx", "XX")],
                  ["Matériel de transport", ("xx", "XX"), "5", ("xx", "XX")],
                  ["Total dotation annuelle", "—", "—", ("xx", "XX")],
              ])

    add_paragraph(doc, "")
    add_paragraph(doc, (
        "Note : Calculer chaque dotation en divisant le montant de l'immobilisation par sa DUP."
    ), bold=True, size=Pt(10))

    # V. Compte d'exploitation prévisionnel
    add_paragraph(doc, "")
    add_subheading(doc, "V. Compte d'exploitation prévisionnel (5 ans)")

    add_table(doc,
              ["Postes", "Année 1", "Année 2", "Année 3", "Année 4", "Année 5"],
              [
                  ["Chiffre d'affaires (CA)", "24 000 000", "30 000 000", "37 500 000", "46 000 000", "56 000 000"],
                  ["Charges variables (CV)", ("xx", "XX"), ("xx", "XX"), ("xx", "XX"), ("xx", "XX"), ("xx", "XX")],
                  ["Eau", ("xx", "XX"), ("xx", "XX"), ("xx", "XX"), ("xx", "XX"), ("xx", "XX")],
                  ["Électricité", ("xx", "XX"), ("xx", "XX"), ("xx", "XX"), ("xx", "XX"), ("xx", "XX")],
                  ["Téléphone", ("xx", "XX"), ("xx", "XX"), ("xx", "XX"), ("xx", "XX"), ("xx", "XX")],
                  ["Carburant", ("xx", "XX"), ("xx", "XX"), ("xx", "XX"), ("xx", "XX"), ("xx", "XX")],
                  ["Loyer", ("xx", "XX"), ("xx", "XX"), ("xx", "XX"), ("xx", "XX"), ("xx", "XX")],
                  ["Assurance", ("xx", "XX"), ("xx", "XX"), ("xx", "XX"), ("xx", "XX"), ("xx", "XX")],
                  ["Autres charges externes", ("xx", "XX"), ("xx", "XX"), ("xx", "XX"), ("xx", "XX"), ("xx", "XX")],
                  ["Total charges externes (CE)", ("xx", "XX"), ("xx", "XX"), ("xx", "XX"), ("xx", "XX"), ("xx", "XX")],
                  ["Charges de personnel (CP)", ("xx", "XX"), ("xx", "XX"), ("xx", "XX"), ("xx", "XX"), ("xx", "XX")],
                  ["Valeur ajoutée (VA = CA − CV − CE)", ("xx", "XX"), ("xx", "XX"), ("xx", "XX"), ("xx", "XX"), ("xx", "XX")],
                  ["EBE (VA − CP)", ("xx", "XX"), ("xx", "XX"), ("xx", "XX"), ("xx", "XX"), ("xx", "XX")],
                  ["Dotations aux amortissements", ("xx", "XX"), ("xx", "XX"), ("xx", "XX"), ("xx", "XX"), ("xx", "XX")],
                  ["Résultat d'exploitation (RE = EBE − Amort.)", ("xx", "XX"), ("xx", "XX"), ("xx", "XX"), ("xx", "XX"), ("xx", "XX")],
                  ["Frais financiers (intérêts)", "910 000", "728 000", "546 000", "364 000", "182 000"],
                  ["Résultat avant impôt (RA = RE − FF)", ("xx", "XX"), ("xx", "XX"), ("xx", "XX"), ("xx", "XX"), ("xx", "XX")],
                  ["IS (30 % × RA)", ("xx", "XX"), ("xx", "XX"), ("xx", "XX"), ("xx", "XX"), ("xx", "XX")],
                  ["Résultat net (RN = RA − IS)", ("xx", "XX"), ("xx", "XX"), ("xx", "XX"), ("xx", "XX"), ("xx", "XX")],
                  ["CAF (RN + Amort.)", ("xx", "XX"), ("xx", "XX"), ("xx", "XX"), ("xx", "XX"), ("xx", "XX")],
              ])

    add_paragraph(doc, "")
    add_paragraph(doc, (
        "Formules de calcul : VA = CA − CV − CE ; EBE = VA − CP ; RE = EBE − Amortissements ; "
        "RA = RE − Frais financiers ; IS = RA × 30 % ; RN = RA − IS ; CAF = RN + Amortissements."
    ), bold=True, size=Pt(10))

    # VI. VAN
    add_paragraph(doc, "")
    add_subheading(doc, "VI. Valeur Actuelle Nette (VAN)")

    add_paragraph(doc, (
        "Taux d'actualisation retenu : 10 %."
    ), bold=True, size=Pt(11))

    add_table(doc,
              ["Année", "CAF (FCFA)", "Coefficient d'actualisation (1/(1+0,10)^n)", "CAF actualisée (FCFA)", "Cumul CAF actualisée (FCFA)"],
              [
                  ["0 (investissement)", "−13 000 000", "1,0000", "−13 000 000", "−13 000 000"],
                  ["1", ("xx", "XX"), "0,9091", ("xx", "XX"), ("xx", "XX")],
                  ["2", ("xx", "XX"), "0,8264", ("xx", "XX"), ("xx", "XX")],
                  ["3", ("xx", "XX"), "0,7513", ("xx", "XX"), ("xx", "XX")],
                  ["4", ("xx", "XX"), "0,6830", ("xx", "XX"), ("xx", "XX")],
                  ["5", ("xx", "XX"), "0,6209", ("xx", "XX"), ("xx", "XX")],
              ])

    add_paragraph(doc, "")
    p = add_paragraph(doc, "")
    run_text(p, "VAN = Cumul des CAF actualisées − Investissement initial = ", bold=True)
    placeholder(p)
    run_text(p, " FCFA")

    add_paragraph(doc, "")
    p = add_paragraph(doc, "")
    run_text(p, "Conclusion : ", bold=True)
    run_text(p, "Si VAN > 0, le projet est rentable et créateur de valeur. Si VAN < 0, le projet "
             "nécessite une révision des hypothèses.")

    # VII. Seuil de rentabilité
    add_paragraph(doc, "")
    add_subheading(doc, "VII. Seuil de rentabilité")

    add_paragraph(doc, "")
    p = add_paragraph(doc, "")
    run_text(p, "Chiffre d'affaires (CA) : ", bold=True)
    run_text(p, "voir tableau compte d'exploitation.")
    p = add_paragraph(doc, "")
    run_text(p, "Charges variables (CV) : ", bold=True)
    placeholder(p)

    p = add_paragraph(doc, "")
    run_text(p, "Marge sur coût variable (MCV) = CA − CV = ", bold=True)
    placeholder(p)

    add_sub2heading(doc, "Charges fixes détaillées")
    add_table(doc,
              ["Poste de charges fixes", "Montant annuel (FCFA)"],
              [
                  ["Loyer", ("xx", "XX")],
                  ["Salaires et charges sociales", ("xx", "XX")],
                  ["Électricité (part fixe)", ("xx", "XX")],
                  ["Téléphone", ("xx", "XX")],
                  ["Eau (part fixe)", ("xx", "XX")],
                  ["Dotations aux amortissements", ("xx", "XX")],
                  ["Frais financiers (intérêts)", "910 000"],
                  ["Total charges fixes (CF)", ("xx", "XX")],
              ])

    add_paragraph(doc, "")
    p = add_paragraph(doc, "")
    run_text(p, "Seuil de rentabilité (SR) = (CA × CF) / MCV = ", bold=True)
    placeholder(p)
    run_text(p, " FCFA")

    p = add_paragraph(doc, "")
    run_text(p, "Point mort (en fraction d'année) = SR / CA = ", bold=True)
    placeholder(p)

    # VIII. Plan de financement 5 ans
    add_paragraph(doc, "")
    add_subheading(doc, "VIII. Plan de financement sur 5 ans")

    add_table(doc,
              ["", "Année 0", "Année 1", "Année 2", "Année 3", "Année 4", "Année 5"],
              [
                  ["RESSOURCES", "", "", "", "", "", ""],
                  ["Apport personnel", "3 900 000", "—", "—", "—", "—", "—"],
                  ["Emprunt bancaire", "9 100 000", "—", "—", "—", "—", "—"],
                  ["CAF", "—", ("xx", "XX"), ("xx", "XX"), ("xx", "XX"), ("xx", "XX"), ("xx", "XX")],
                  ["Subventions", ("xx", "XX"), ("xx", "XX"), ("xx", "XX"), ("xx", "XX"), ("xx", "XX"), ("xx", "XX")],
                  ["Cessions d'actifs", "—", "—", "—", "—", "—", "—"],
                  ["Total ressources", "13 000 000", ("xx", "XX"), ("xx", "XX"), ("xx", "XX"), ("xx", "XX"), ("xx", "XX")],
                  ["", "", "", "", "", "", ""],
                  ["EMPLOIS", "", "", "", "", "", ""],
                  ["Investissements", "10 650 000", "—", "—", ("xx", "XX"), "—", ("xx", "XX")],
                  ["BFR initial", "2 350 000", "—", "—", "—", "—", "—"],
                  ["Variation du BFR", "—", ("xx", "XX"), ("xx", "XX"), ("xx", "XX"), ("xx", "XX"), ("xx", "XX")],
                  ["Remboursement emprunt (EP)", "—", "1 820 000", "1 820 000", "1 820 000", "1 820 000", "1 820 000"],
                  ["Total emplois", "13 000 000", ("xx", "XX"), ("xx", "XX"), ("xx", "XX"), ("xx", "XX"), ("xx", "XX")],
                  ["", "", "", "", "", "", ""],
                  ["Trésorerie nette (Ressources − Emplois)", "0", ("xx", "XX"), ("xx", "XX"), ("xx", "XX"), ("xx", "XX"), ("xx", "XX")],
                  ["Trésorerie cumulée", "0", ("xx", "XX"), ("xx", "XX"), ("xx", "XX"), ("xx", "XX"), ("xx", "XX")],
              ])

    add_page_break(doc)


# ──────────────── 5ème PARTIE ────────────────

def build_partie5(doc):
    add_heading_custom(doc, "5ème PARTIE : IMPACTS SOCIO-ÉCONOMIQUES ET GESTION DES RISQUES",
                       size=Pt(14), alignment=WD_ALIGN_PARAGRAPH.CENTER)
    add_paragraph(doc, "")

    # I. Impact économique
    add_subheading(doc, "I. Impact économique")

    p = add_paragraph(doc, "")
    run_text(p, "Masse salariale annuelle : ", bold=True)
    placeholder(p)
    run_text(p, " FCFA")

    add_bullet(doc, "Contribution fiscale : IS, TVA, Patente, CFCE, IMF")
    add_bullet(doc, "Achat de matières premières locales (fruits, légumes, emballages)")
    add_bullet(doc, "Chiffre d'affaires prévisionnel de 56 000 000 FCFA en année 5")
    add_bullet(doc, "Dynamisation de la filière agroalimentaire dans la région de Kaffrine")

    # II. Impact social
    add_subheading(doc, "II. Impact social")
    add_bullet(doc, "Lutte contre le chômage : création de 7 à 17 emplois directs, priorité aux jeunes et aux femmes")
    add_bullet(doc, "Réduction de la pauvreté : revenus stables pour les employés et leurs familles")
    add_bullet(doc, "Sécurité alimentaire : produits frais et nutritifs disponibles toute l'année")
    add_bullet(doc, "Revalorisation sociale : promotion de l'entrepreneuriat local")
    add_bullet(doc, "Amélioration des conditions de vie : aliments sains et abordables")
    add_bullet(doc, "Formation et transfert de compétences en agriculture moderne")
    add_bullet(doc, "Promotion du rôle des femmes dans l'économie locale")

    # III. Impact environnemental
    add_subheading(doc, "III. Impact environnemental")
    add_bullet(doc, "Économie d'eau de 90 % grâce au système hydroponique en circuit fermé")
    add_bullet(doc, "Production locale réduisant les émissions liées au transport")
    add_bullet(doc, "Absence de pesticides chimiques (production sous serre contrôlée)")
    add_bullet(doc, "Compostage des déchets organiques pour amendement des sols")
    add_bullet(doc, "Préservation des sols (culture hors sol)")
    add_bullet(doc, "Recyclage de l'eau en circuit fermé")
    add_bullet(doc, "Objectif à moyen terme : passage à l'énergie solaire")

    # IV. Risques et mesures
    add_subheading(doc, "IV. Risques et mesures d'atténuation")

    add_table(doc,
              ["N°", "Risque", "Probabilité", "Impact", "Mesure d'atténuation"],
              [
                  ["1", "Coupures d'électricité", "Élevée", "Élevé", "Groupe électrogène + transition solaire"],
                  ["2", "Pénurie d'eau", "Moyenne", "Élevé", "Forage/puits + citernes + recyclage circuit fermé"],
                  ["3", "Panne technique", "Moyenne", "Élevé", "Formation + stock de pièces de rechange + maintenance préventive"],
                  ["4", "Mévente", "Faible", "Élevé", "Diversification des canaux + politique de prix + promotion active"],
                  ["5", "Hausse du prix des intrants", "Moyenne", "Moyen", "Achat groupé + contrats fournisseurs + substituts locaux"],
                  ["6", "Concurrence accrue", "Faible", "Moyen", "Innovation produit + fidélisation client + qualité supérieure"],
                  ["7", "Chaleur extrême", "Moyenne", "Moyen", "Serres ventilées + brumisation + variétés résistantes"],
                  ["8", "Difficultés de remboursement", "Faible", "Élevé", "Réserve de trésorerie + gestion financière rigoureuse"],
                  ["9", "Risques sanitaires", "Faible", "Élevé", "Protocole HACCP + contrôle qualité + traçabilité"],
                  ["10", "Non-conformité réglementaire", "Faible", "Moyen", "Veille réglementaire + accompagnement juridique"],
              ])

    add_page_break(doc)


# ──────────────── CONCLUSION ────────────────

def build_conclusion(doc):
    add_heading_custom(doc, "CONCLUSION", size=Pt(14),
                       alignment=WD_ALIGN_PARAGRAPH.CENTER)
    add_paragraph(doc, "")

    add_paragraph(doc, (
        "Le présent projet d'agriculture hydroponique et de transformation agroalimentaire "
        "répond à un besoin réel et croissant en produits frais et sains dans la région de Kaffrine. "
        "Avec un investissement total de 13 000 000 FCFA, le projet prévoit un chiffre d'affaires "
        "évoluant de 24 000 000 FCFA en année 1 à 56 000 000 FCFA en année 5."
    ), first_line_indent=Cm(1.27))

    add_paragraph(doc, (
        "Au-delà de sa dimension économique, ce projet porte des ambitions sociales fortes : "
        "création d'emplois durables pour les jeunes et les femmes, contribution à la sécurité "
        "alimentaire, promotion d'une agriculture moderne et durable. Les risques identifiés "
        "sont maîtrisables grâce à des mesures d'atténuation concrètes."
    ), first_line_indent=Cm(1.27))

    add_paragraph(doc, (
        "Les perspectives de développement sont prometteuses : introduction de l'aquaponie, "
        "extension des capacités de production, développement de nouveaux produits transformés "
        "et expansion vers le marché national. Ce business plan constitue une base solide pour "
        "la mobilisation des financements nécessaires à la concrétisation de ce projet."
    ), first_line_indent=Cm(1.27))

    add_page_break(doc)


# ──────────────── ANNEXES ────────────────

def build_annexes(doc):
    add_heading_custom(doc, "ANNEXES", size=Pt(14),
                       alignment=WD_ALIGN_PARAGRAPH.CENTER)
    add_paragraph(doc, "")

    # Annexe 1
    add_subheading(doc, "Annexe 1 : Questionnaire d'enquête terrain")
    add_paragraph(doc, "")

    questions = [
        ("1.", "Achetez-vous régulièrement des légumes frais ?", "☐ Oui   ☐ Non"),
        ("2.", "Où achetez-vous habituellement vos légumes ?", "☐ Marché   ☐ Boutique   ☐ Producteur   ☐ Autre"),
        ("3.", "Quel est votre budget hebdomadaire en légumes ?", "☐ < 2 000 FCFA   ☐ 2 000 – 5 000   ☐ 5 000 – 10 000   ☐ > 10 000 FCFA"),
        ("4.", "Seriez-vous intéressé(e) par des légumes cultivés localement sous serre ?", "☐ Oui   ☐ Non   ☐ Peut-être"),
        ("5.", "Achetez-vous des produits transformés (jus, confitures) ?", "☐ Oui   ☐ Non"),
        ("6.", "Quel prix accepteriez-vous pour un jus naturel local de 33 cl ?", "☐ < 500 FCFA   ☐ 500 – 1 000   ☐ 1 000 – 1 500   ☐ > 1 500 FCFA"),
        ("7.", "Quels critères sont les plus importants pour vous ?", "☐ Prix   ☐ Qualité   ☐ Origine locale   ☐ Disponibilité"),
        ("8.", "Connaissez-vous l'hydroponie (culture hors sol) ?", "☐ Oui   ☐ Non"),
    ]

    for num, q, opts in questions:
        p = add_paragraph(doc, "", space_after=Pt(2))
        run_text(p, f"{num} ", bold=True)
        run_text(p, q)
        add_paragraph(doc, opts, space_after=Pt(8))

    add_paragraph(doc, "")

    # Annexe 2
    add_subheading(doc, "Annexe 2 : Tableau des produits et prix prévisionnels")
    add_paragraph(doc, "")

    add_table(doc,
              ["Produit", "Unité", "Prix unitaire (FCFA)", "Volume mensuel estimé"],
              [
                  ["Tomate", "kg", ("xx", "XX"), ("xx", "XX")],
                  ["Laitue", "pièce", ("xx", "XX"), ("xx", "XX")],
                  ["Poivron", "kg", ("xx", "XX"), ("xx", "XX")],
                  ["Concombre", "kg", ("xx", "XX"), ("xx", "XX")],
                  ["Aubergine", "kg", ("xx", "XX"), ("xx", "XX")],
                  ["Jus de fruit (33 cl)", "bouteille", ("xx", "XX"), ("xx", "XX")],
                  ["Confiture (250 g)", "pot", ("xx", "XX"), ("xx", "XX")],
                  ["Sauce tomate (500 ml)", "bouteille", ("xx", "XX"), ("xx", "XX")],
                  ["Compote (250 g)", "pot", ("xx", "XX"), ("xx", "XX")],
                  ["Tartinade (200 g)", "pot", ("xx", "XX"), ("xx", "XX")],
              ])


# ──────────────── MAIN ────────────────

def main():
    doc = Document()

    # --- Page setup ---
    for section in doc.sections:
        section.top_margin = Cm(2.5)
        section.bottom_margin = Cm(2.5)
        section.left_margin = Cm(2.5)
        section.right_margin = Cm(2.5)

    # --- Default style ---
    style = doc.styles['Normal']
    style.font.name = FONT_NAME
    style.font.size = Pt(12)
    style.font.color.rgb = BLACK
    rPr = style.element.get_or_add_rPr()
    rFonts = parse_xml(f'<w:rFonts {nsdecls("w")} w:ascii="{FONT_NAME}" w:hAnsi="{FONT_NAME}" w:eastAsia="{FONT_NAME}" w:cs="{FONT_NAME}"/>')
    rPr.insert(0, rFonts)
    style.paragraph_format.line_spacing = Pt(18)

    # --- Build sections ---
    build_cover_page(doc)
    build_toc(doc)
    build_introduction(doc)
    build_partie1(doc)
    build_partie2(doc)
    build_partie3(doc)
    build_partie4(doc)
    build_partie5(doc)
    build_conclusion(doc)
    build_annexes(doc)

    # --- Page numbers ---
    add_page_numbers(doc)

    # --- Save ---
    doc.save(OUTPUT)
    print(f"✓ Document généré : {OUTPUT}")


if __name__ == "__main__":
    main()
